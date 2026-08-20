# -*- coding: utf-8 -*-
"""
欠席チェック君 (GUI版)

やること:
  1. PDF「欠席時対応記録票」をウィンドウにドラッグ&ドロップ(またはボタンで選択)
  2. カレンダー画面からコピーしたテキストを、下のテキスト欄に貼り付け(Ctrl+V)
  3. 「比較する」ボタンを押すと、一致件数・不一致(片方にしかない組み合わせ)を
     ウィンドウ内に一覧表示する。

■事前準備 (初回のみ、コマンドプロンプトで実行)
    pip install pdfplumber tkinterdnd2

■起動方法
    python compare_absences_gui.py

  ドラッグ&ドロップができない場合は自動的に「PDFを選択」ボタンのみのモードで動作します。
"""

import json
import os
import re
import sys
import unicodedata
import webbrowser
import calendar as calendar_mod
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter import font as tkfont
import openpyxl

# ============================== 仕事の進め方 / 取説 ==============================

INSTRUCTIONS_FILENAME = "手順_欠席加算比較チェック君.md"

DEFAULT_INSTRUCTIONS = """\
※この手順は変更されることがあります。内容を直接書き換えたい場合は
　右上の「✎ 取説を編集」ボタンを押してください(メモ帳でこのファイルが開きます)。
※URL(https://...)やファイルのパス(例: C:\\folder\\file.xlsx)をそのまま1行に
　書いておくと、「仕事の進め方」画面でクリックできるリンクになります。
"""

_LINK_CHAR = r"[^\s。、，」』（）\"']"
LINK_PATTERN = re.compile(rf"(https?://{_LINK_CHAR}+|[A-Za-z]:\\{_LINK_CHAR}+|\\\\{_LINK_CHAR}+)")
LINK_TRAILING_CHARS = "。、,.)）」』\"'"


def get_instructions_path():
    app_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(app_dir, INSTRUCTIONS_FILENAME)


def ensure_instructions_file():
    path = get_instructions_path()
    if not os.path.exists(path):
        with open(path, "w", encoding="utf-8") as f:
            f.write(DEFAULT_INSTRUCTIONS)
    return path


def read_instructions():
    path = ensure_instructions_file()
    try:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    except OSError as e:
        return f"手順ファイルの読み込みに失敗しました。\n{path}\n{e}"


# ============================== 前回設定の保存(対象日付範囲・結果欄の高さ) ==============================

SETTINGS_FILENAME = "gui_settings.json"


def get_settings_path():
    app_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(app_dir, SETTINGS_FILENAME)


def load_settings():
    path = get_settings_path()
    if not os.path.exists(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_settings(settings):
    path = get_settings_path()
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)
    except OSError:
        pass


try:
    import pdfplumber
except ImportError:
    pdfplumber = None

# tkinterdnd2があればドラッグ&ドロップ対応、無ければ通常のtkinterにフォールバック
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    DND_AVAILABLE = True
except ImportError:
    DND_AVAILABLE = False


# ============ 施設検索(店舗検索/tenpo.pyと同じ考え方) ============

STORE_EXCEL_PATH = r"C:\Users\owner\Desktop\works\保管書類\店舗リスト\table_HHHL共通店舗一覧m.xlsm"
STORE_SHEET_NAME = "リスト"
STORE_NAME_COL = "店舗名"
STORE_AREA_COL = "区分"


class StoreData:
    def __init__(self, path, sheet_name, name_col, area_col):
        self.path = path
        self.sheet_name = sheet_name
        self.name_col = name_col
        self.area_col = area_col
        self.rows = []  # [(店舗名, エリア), ...]
        self.error = None
        self.reload()

    def reload(self):
        try:
            wb = openpyxl.load_workbook(
                self.path, read_only=True, data_only=True, keep_vba=True
            )
            ws = wb[self.sheet_name]
            header = None
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    header = row
                    try:
                        name_idx = header.index(self.name_col)
                        area_idx = header.index(self.area_col)
                    except ValueError:
                        self.error = "見出し行に「{}」または「{}」が見つかりません".format(
                            self.name_col, self.area_col
                        )
                        self.rows = []
                        wb.close()
                        return
                    continue
                name = row[name_idx] if name_idx < len(row) else None
                area = row[area_idx] if area_idx < len(row) else None
                if name:
                    rows.append((str(name), str(area) if area else ""))
            wb.close()
            self.rows = rows
            self.error = None
        except FileNotFoundError:
            self.error = "エクセルファイルが見つかりません:\n{}".format(self.path)
            self.rows = []
        except PermissionError:
            self.error = "エクセルファイルを開けません（他で使用中の可能性があります）"
            self.rows = []
        except Exception as e:
            self.error = "読み込みエラー: {}".format(e)
            self.rows = []

    def search(self, keyword):
        if not keyword:
            return []
        kw = keyword.casefold()
        return [
            (name, area)
            for name, area in self.rows
            if kw in name.casefold()
        ]


def extract_facility_name(text: str):
    """HUGカレンダー画面の貼り付けテキストから、「施設」という見出し行の
    すぐ下にある地名(例:「岩曽」)を取り出す。見つからなければNoneを返す。"""
    lines = text.splitlines()
    for i, line in enumerate(lines):
        if line.strip() == "施設":
            for nxt in lines[i + 1:]:
                s = nxt.strip()
                if s:
                    return s
            break
    return None


# ============ 突き合わせロジック(compare_absences.pyと同じ考え方) ============

RELATION_WORDS = ["祖父母", "祖父", "祖母", "本人", "その他", "父", "母", "兄", "姉", "弟", "妹"]


# 同一人物でもPDF/カレンダーで表記が揺れる異体字を吸収するための対応表
# (例: 「田邉」と「田邊」は同じ「たなべ」姓として一致させる)
VARIANT_CHAR_MAP = {
    "邊": "邉",
}


def normalize_name(name: str) -> str:
    if name is None:
        return ""
    name = unicodedata.normalize("NFKC", name)
    name = name.upper()
    name = re.sub(r"\s+", "", name)
    for src, dst in VARIANT_CHAR_MAP.items():
        name = name.replace(src, dst)
    return name


def extract_day(date_str: str):
    """日付文字列から末尾の"日"の数字だけを取り出す('7/16'→16, '16'→16)"""
    if date_str is None:
        return None
    s = unicodedata.normalize("NFKC", date_str).strip()
    m = re.search(r"(\d{1,2})\s*$", s)
    return int(m.group(1)) if m else None


def format_mmdd(date_str: str, month=None) -> str:
    """
    表示用に日付を 'MM/DD'(ゼロ埋め2桁) 形式へ整える。
    date_strに月が含まれていれば(例:'7/1')それを使い、
    月が無ければ(例:'1'だけ)引数のmonthを補って使う。
    """
    if date_str is None:
        return ""
    s = unicodedata.normalize("NFKC", date_str).strip()
    m = re.search(r"(?:(\d{1,2})[/\-])?(\d{1,2})\s*$", s)
    if not m:
        return s
    mm = m.group(1)
    dd = m.group(2)
    if mm is None:
        mm = month
    if mm is None:
        return f"?/{int(dd):02d}"
    return f"{int(mm):02d}/{int(dd):02d}"


def normalize_date(date_str: str) -> str:
    """
    同じ月のPDFとカレンダーを比較する前提のため、月の有無によらず
    末尾の"日"の数字だけを取り出して正規化する('7/1'も'1'も同じキーになる)。
    """
    if date_str is None:
        return ""
    date_str = unicodedata.normalize("NFKC", date_str).strip()
    m = re.search(r"(\d{1,2})\s*$", date_str)
    if not m:
        return date_str
    return str(int(m.group(1)))


def _extract_pdf_table_rows(pdf_path: str):
    """
    PDFの表を1行ずつパースして、
    [{'row_no':..., 'flag':..., 'name':..., 'date':(欠席日), 'uketsuke':(受付日),
      'has_blank':(記録票の必須項目に空欄があるか)}, ...]
    を返す(○が付いているかどうかに関わらず全行)。
    """
    if pdfplumber is None:
        raise RuntimeError("pdfplumberがインストールされていません。\npip install pdfplumber を実行してください。")

    rows_out = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                header = table[0]

                def find_col(keyword):
                    for i, h in enumerate(header):
                        if h and keyword in h.replace("\n", ""):
                            return i
                    return None

                def find_col_exact(keyword):
                    # "次回利用日"と"次回利用日の確認"のように、部分一致だと
                    # 別の列を誤って拾ってしまう場合があるための完全一致版
                    for i, h in enumerate(header):
                        if h and h.replace("\n", "") == keyword:
                            return i
                    return None

                idx_flag = find_col("算定")
                idx_name = find_col("児童氏名")
                idx_rel = find_col("続柄")
                idx_date = find_col("欠席日")
                idx_uketsuke = find_col("受付日")
                if idx_flag is None or idx_name is None or idx_date is None:
                    continue

                # 欠席加算を算定する場合に記入必須となる項目(ルール③)
                idx_reason = find_col("欠席理由")
                idx_child_state = find_col("児童の様子")
                idx_support = find_col("援助内容")
                idx_next_confirm = find_col_exact("次回利用日の確認")
                idx_next_date = find_col_exact("次回利用日")
                idx_recorder = find_col("記入者")

                # 「欠席加算」列の左隣にある通し番号列(ヘッダーは空欄)を行番号として使う
                idx_rowno = idx_flag - 1 if idx_flag is not None and idx_flag > 0 else None

                for row in table[1:]:
                    if not row or len(row) <= max(idx_flag, idx_name, idx_date):
                        continue
                    flag = (row[idx_flag] or "").strip()
                    name = (row[idx_name] or "").strip()
                    date = (row[idx_date] or "").strip()
                    uketsuke = (row[idx_uketsuke] or "").strip() if idx_uketsuke is not None else ""
                    rel_cell = (row[idx_rel] or "").strip() if idx_rel is not None else ""
                    row_no = (row[idx_rowno] or "").strip() if idx_rowno is not None else ""

                    final_relation = rel_cell
                    if rel_cell and rel_cell not in RELATION_WORDS:
                        # 氏名が長くて続柄列にはみ出しているケースを補正
                        combined = name + rel_cell
                        final_relation = ""
                        for w in RELATION_WORDS:
                            if combined.endswith(w):
                                name = combined[: -len(w)]
                                final_relation = w
                                break

                    def cell(idx):
                        if idx is None or idx >= len(row):
                            return ""
                        return (row[idx] or "").strip()

                    required_values = [
                        uketsuke,
                        final_relation,
                        cell(idx_reason),
                        cell(idx_child_state),
                        cell(idx_support),
                        cell(idx_next_confirm),
                        cell(idx_next_date),
                        cell(idx_recorder),
                    ]
                    has_blank = any(v == "" for v in required_values)

                    if name and date:
                        rows_out.append(
                            {
                                "row_no": row_no,
                                "flag": flag,
                                "name": name,
                                "date": date,
                                "uketsuke": uketsuke,
                                "has_blank": has_blank,
                            }
                        )
    return rows_out


def is_maru(flag: str) -> bool:
    """
    欠席加算(算定)欄の判定。「○」はもちろん、空欄(未記入)の場合も
    ○が入力されているものとして扱う。
    """
    return (flag or "").strip() in ("○", "")


def extract_pdf_rows(pdf_path: str):
    """PDFから「欠席加算=○(空欄含む)」の行を [{'date':..., 'name':..., 'row_no':..., 'has_blank':...}, ...] で返す"""
    all_rows = _extract_pdf_table_rows(pdf_path)
    return [
        {"date": r["date"], "name": r["name"], "row_no": r["row_no"], "has_blank": r["has_blank"]}
        for r in all_rows
        if is_maru(r["flag"])
    ]


def find_pdf_blank_flag_rows(pdf_path: str):
    """PDFの「欠席加算」欄が空欄(未記入)になっている行を返す。
    空欄は計算上○として扱われるが、記入漏れ確認用に一覧化する。
    [{'name':..., 'date':..., 'row_no':...}, ...]"""
    all_rows = _extract_pdf_table_rows(pdf_path)
    return [r for r in all_rows if (r["flag"] or "").strip() == ""]


def get_pdf_flag_map(pdf_path: str):
    """(欠席日, 氏名)ごとの「欠席加算」欄の記載内容とPDF行番号を返す(○以外も含む全行対象)。
    only_cal(HUGにはあるが日報側に○が無い)の表示で、本当に記載が無いのか
    「×」等が明示的に書いてあるのかを区別し、行番号も表示するために使う。
    戻り値: {(date_key, name_key): {'flag':..., 'row_no':...}, ...}"""
    all_rows = _extract_pdf_table_rows(pdf_path)
    result = {}
    for r in all_rows:
        key = (normalize_date(r["date"]), normalize_name(r["name"]))
        result[key] = {"flag": r["flag"], "row_no": r["row_no"]}
    return result


def find_pdf_duplicates(pdf_path: str):
    """PDF内で「氏名+欠席日」が完全一致する行(=重複入力の疑い)をまとめて返す。
    ○以外の行も含めて全行を対象にする。
    [{'name':..., 'date':..., 'row_nos': [...]}, ...]"""
    all_rows = _extract_pdf_table_rows(pdf_path)
    groups = {}
    for r in all_rows:
        key = (normalize_date(r["date"]), normalize_name(r["name"]))
        groups.setdefault(key, []).append(r)

    duplicates = []
    for rows in groups.values():
        if len(rows) > 1:
            duplicates.append(
                {
                    "name": rows[0]["name"],
                    "date": rows[0]["date"],
                    "row_nos": [r["row_no"] for r in rows],
                }
            )
    return duplicates


def parse_month_day(date_str: str):
    """'7/1'のような文字列を(月,日)のタプルにする。読み取れなければNone"""
    if not date_str:
        return None
    s = unicodedata.normalize("NFKC", date_str).strip()
    m = re.search(r"(\d{1,2})[/\-](\d{1,2})", s)
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def find_pdf_date_errors(pdf_path: str):
    """
    PDFの「欠席加算=○(空欄含む)」の行だけを対象に、受付日が欠席日より後になっている
    (＝欠席日 < 受付日、日付の誤記の疑いがある)行を
    [{'row_no':..., 'name':..., 'date':..., 'uketsuke':...}, ...] で返す。
    """
    all_rows = _extract_pdf_table_rows(pdf_path)
    errors = []
    for r in all_rows:
        if not is_maru(r["flag"]):
            continue
        d1 = parse_month_day(r["date"])       # 欠席日
        d2 = parse_month_day(r["uketsuke"])    # 受付日
        if d1 is not None and d2 is not None and d1 < d2:
            errors.append(r)
    return errors


def parse_calendar_text(text: str, month: int = None):
    """カレンダー画面のコピペテキストから「欠席」を抽出。
    通常の「欠席」はstatus='○'、「欠席（加算なし）」はstatus='✕'として区別する。"""
    lines = text.splitlines()

    def is_day_header(line):
        m = re.fullmatch(r"(?:\S+\s+)?(\d{1,2})", line.strip())
        if m:
            n = int(m.group(1))
            if 1 <= n <= 31:
                return n
        return None

    results = []
    current_day = None
    state = None

    for line in lines:
        s = line.strip()
        if not s:
            continue
        if s.startswith("日\t月") or s.startswith("日 月"):
            continue

        dnum = is_day_header(s)
        if dnum is not None:
            current_day = dnum
            state = None
            continue

        if s == "休み":
            state = "skip"
            continue
        if re.fullmatch(r"出席(\d+)人", s):
            state = "attend"
            continue
        if re.fullmatch(r"欠席（加算なし）(\d+)人", s):
            state = "absent1_nokasan"
            continue
        if re.fullmatch(r"欠席(\d+)人", s):
            state = "absent1"
            continue

        if current_day is not None and state in ("absent1", "absent1_nokasan"):
            date_label = f"{month}/{current_day}" if month else str(current_day)
            status = "○" if state == "absent1" else "✕"
            results.append({"date": date_label, "name": s, "status": status})

    return results


def compare_rows(pdf_rows, cal_rows):
    def key(r):
        return (normalize_date(r["date"]), normalize_name(r["name"]))

    pdf_map = {key(r): r for r in pdf_rows}
    cal_map = {key(r): r for r in cal_rows}

    pdf_keys, cal_keys = set(pdf_map), set(cal_map)
    matched = pdf_keys & cal_keys
    only_pdf = sorted(pdf_keys - cal_keys)
    only_cal = sorted(cal_keys - pdf_keys)

    matched_pdf_rows = [pdf_map[k] for k in sorted(matched)]

    return matched_pdf_rows, [pdf_map[k] for k in only_pdf], [cal_map[k] for k in only_cal]


# ============================== GUI ==============================

BaseTk = TkinterDnD.Tk if DND_AVAILABLE else tk.Tk


class App(BaseTk):
    def __init__(self):
        super().__init__()
        self.title("欠席チェック君")
        self.geometry("820x720")

        self.pdf_path = None
        self._instructions_win = None
        self.settings = load_settings()
        self.store_data = StoreData(STORE_EXCEL_PATH, STORE_SHEET_NAME, STORE_NAME_COL, STORE_AREA_COL)

        self._build_widgets()
        self._refresh_store_status()

        # メイン画面が表示された直後に、仕事の進め方ガイドを自動表示する
        self.after(200, lambda: self._show_instructions_dialog(auto=True))

    def _build_widgets(self):
        pad = {"padx": 10, "pady": 6}

        header_row = ttk.Frame(self)
        header_row.pack(side="top", fill="x", padx=10, pady=(8, 0))
        ttk.Button(header_row, text="📋 仕事の進め方", command=self._show_instructions_dialog).pack(side="right")
        ttk.Button(header_row, text="✎ 取説を編集", command=self._open_instructions_editor).pack(side="right", padx=(0, 6))

        # --- 上段(①②、施設検索まで)は必要最小限の高さで固定。サッシは廃止し、
        #     下段(結果)だけがウィンドウの余った高さを埋める ---
        top_frame = ttk.Frame(self)
        top_frame.pack(side="top", fill="x")

        bottom_frame = ttk.Frame(self)
        bottom_frame.pack(side="top", fill="both", expand=True)

        # --- PDFエリア ---
        frame_pdf = ttk.LabelFrame(top_frame, text="① 欠席時対応記録票(日報PDF)")
        frame_pdf.pack(fill="x", **pad)

        drop_text = "ここにPDFをドラッグ&ドロップ" if DND_AVAILABLE else "PDFファイルを選択してください"
        self.pdf_label = tk.Label(
            frame_pdf, text=drop_text, bg="#f0f0f0", height=3, relief="groove"
        )
        self.pdf_label.pack(fill="x", padx=10, pady=(10, 4))

        if DND_AVAILABLE:
            self.pdf_label.drop_target_register(DND_FILES)
            self.pdf_label.dnd_bind("<<Drop>>", self._on_pdf_drop)

        pdf_button_row = ttk.Frame(frame_pdf)
        pdf_button_row.pack(pady=(0, 10))
        ttk.Button(pdf_button_row, text="PDFを選択...", command=self._choose_pdf).pack(side="left")
        ttk.Button(pdf_button_row, text="開く", command=self._open_pdf).pack(side="left", padx=(6, 0))

        # --- ③ボタン〜施設検索〜検索結果窓をまとめて、常に現在の最小の高さで固定する ---
        bottom_fixed_area = tk.Frame(top_frame)
        bottom_fixed_area.pack(side="bottom", fill="x")

        ttk.Button(bottom_fixed_area, text="③ 比較する", command=self._run_compare).pack(pady=6)

        # --- 施設検索(②に貼り付けたテキストの「施設」の下の地名を自動転記) ---
        search_area = ttk.Frame(bottom_fixed_area)
        search_area.pack(fill="x")

        search_row = ttk.Frame(search_area)
        search_row.pack(side="top", fill="x", padx=10, pady=(8, 0))
        ttk.Label(search_row, text="施設検索:").pack(side="left")
        self.store_entry = ttk.Entry(search_row, width=16)
        self.store_entry.pack(side="left", padx=(4, 4))
        self.store_entry.bind("<KeyRelease>", self._on_store_search)
        ttk.Button(search_row, text="更新", width=4, command=self._reload_store_data).pack(side="left")
        self.store_status_var = tk.StringVar()
        ttk.Label(search_row, textvariable=self.store_status_var, foreground="#666666").pack(
            side="left", padx=(10, 0)
        )

        store_result_frame = ttk.Frame(search_area)
        store_result_frame.pack(side="top", anchor="w", padx=10, pady=(2, 4))
        self.store_result_text = tk.Text(store_result_frame, height=3, width=40, wrap="none")
        store_scroll = ttk.Scrollbar(
            store_result_frame, orient="vertical", command=self.store_result_text.yview
        )
        self.store_result_text.config(yscrollcommand=store_scroll.set, state="disabled")
        store_scroll.pack(side="right", fill="y")
        self.store_result_text.pack(side="left", fill="y")
        self._add_copy_context_menu(self.store_result_text)

        # --- 現時点の必要最小の高さを測って固定し、以後どんなにドラッグ・リサイズしても変わらないようにする ---
        self.update_idletasks()
        bottom_fixed_area.configure(height=bottom_fixed_area.winfo_reqheight())
        bottom_fixed_area.pack_propagate(False)

        # --- ②の上にHUG側の画面場所を示すヒント行(他より1段階大きいフォント) ---
        default_font = tkfont.nametofont("TkDefaultFont")
        point_font = tkfont.Font(
            family=default_font.cget("family"), size=default_font.cget("size") + 2
        )
        ttk.Label(
            top_frame,
            text="Point【MENU-利用者管理-出席カレンダー】",
            font=point_font,
            anchor="w",
        ).pack(fill="x", padx=10, pady=(6, 0))

        # --- カレンダーテキストエリア ---
        frame_cal = ttk.LabelFrame(top_frame, text="② HUGカレンダー画面のテキストを貼り付け(全選択コピペでOK)")
        frame_cal.pack(fill="both", expand=True, **pad)

        top_row = ttk.Frame(frame_cal)
        top_row.pack(fill="x", padx=10, pady=(10, 0))
        tk.Label(top_row, text="今は何年何月ですか").pack(side="left")
        self.year_entry = ttk.Entry(top_row, width=6)
        self.year_entry.pack(side="left", padx=(6, 2))
        self.year_entry.insert(0, str(datetime.now().year))
        tk.Label(top_row, text="年").pack(side="left")
        self.month_entry = ttk.Entry(top_row, width=4)
        self.month_entry.pack(side="left", padx=(6, 2))
        self.month_entry.insert(0, str(datetime.now().month))
        tk.Label(top_row, text="月　(結果はMM/DD形式で表示されます)").pack(side="left")

        range_row = ttk.Frame(frame_cal)
        range_row.pack(fill="x", padx=10, pady=(6, 0))
        range_font = tkfont.Font(
            family=default_font.cget("family"), size=default_font.cget("size") + 4, weight="bold"
        )
        range_font_selected = tkfont.Font(
            family=default_font.cget("family"), size=default_font.cget("size") + 6, weight="bold"
        )
        range_style = ttk.Style()
        range_style.configure("Range.TRadiobutton", font=range_font, foreground="#000000")
        range_style.configure("RangeSelected.TRadiobutton", font=range_font_selected, foreground="#c62828")
        tk.Label(range_row, text="対象日付範囲:", font=range_font).pack(side="left")
        valid_ranges = ("first_third", "second_third", "third_third")
        saved_range = self.settings.get("range")
        self.range_var = tk.StringVar(value=saved_range if saved_range in valid_ranges else "first_third")
        self.range_radios = {}
        for text, value, padx in (
            ("1日〜15日", "first_third", (6, 0)),
            ("16日〜25日", "second_third", (10, 0)),
            ("26日〜月末", "third_third", (10, 0)),
        ):
            radio = ttk.Radiobutton(
                range_row,
                text=text,
                variable=self.range_var,
                value=value,
                style="Range.TRadiobutton",
            )
            radio.pack(side="left", padx=padx)
            self.range_radios[value] = radio

        def _update_range_emphasis(*_args):
            selected = self.range_var.get()
            for value, radio in self.range_radios.items():
                radio.configure(
                    style="RangeSelected.TRadiobutton" if value == selected else "Range.TRadiobutton"
                )
            self._save_settings()

        self.range_var.trace_add("write", _update_range_emphasis)
        _update_range_emphasis()

        cal_text_frame = ttk.Frame(frame_cal)
        cal_text_frame.pack(fill="x", padx=10, pady=10)
        self.cal_text = tk.Text(cal_text_frame, height=2, wrap="none")
        cal_scroll = ttk.Scrollbar(cal_text_frame, orient="vertical", command=self.cal_text.yview)
        self.cal_text.config(yscrollcommand=cal_scroll.set)
        cal_scroll.pack(side="right", fill="y")
        self.cal_text.pack(side="left", fill="both", expand=True)
        self._add_edit_context_menu(self.cal_text)
        self.cal_text.bind("<<Paste>>", self._on_cal_text_paste)

        # --- 結果表示エリア(左右にドラッグで幅調整、下端のつまみで左右2つの高さを同時に調整可能) ---
        frame_result = ttk.LabelFrame(bottom_frame, text="結果")
        frame_result.pack(fill="both", expand=True, **pad)

        paned = tk.PanedWindow(frame_result, orient="horizontal", sashrelief="raised", sashwidth=6)
        paned.pack(fill="both", expand=True, padx=10, pady=(10, 4))

        left_frame, self.result_text = self._make_result_box(paned, "不一致・日付誤記の疑い")
        paned.add(left_frame, stretch="always")

        right_frame, self.blank_text = self._make_result_box(paned, "一致しているが〇の行に空欄がある")
        paned.add(right_frame, stretch="always")

        # --- 下端のつまみをドラッグすると、左右2つの結果欄の高さが同時に変わる ---
        grip = tk.Frame(frame_result, height=6, bg="#cfcfcf", cursor="sb_v_double_arrow")
        grip.pack(side="bottom", fill="x", padx=10, pady=(0, 4), before=paned)
        drag_state = {"y0": 0, "height0": 0}

        def on_press(event):
            drag_state["y0"] = event.y_root
            drag_state["height0"] = int(self.result_text.cget("height"))

        def on_drag(event):
            info = self.result_text.dlineinfo("1.0")
            line_h = info[3] if info else 16
            delta_lines = int(round((event.y_root - drag_state["y0"]) / line_h))
            new_height = max(4, drag_state["height0"] + delta_lines)
            if new_height != int(self.result_text.cget("height")):
                self.result_text.config(height=new_height)
                self.blank_text.config(height=new_height)

        def on_release(_event):
            self._save_settings()

        grip.bind("<Button-1>", on_press)
        grip.bind("<B1-Motion>", on_drag)
        grip.bind("<ButtonRelease-1>", on_release)

    # --- 結果欄のテキスト枠を1つ作成する ---
    def _make_result_box(self, parent, title):
        frame = ttk.Frame(parent)
        ttk.Label(frame, text=title).pack(anchor="w")

        text_frame = ttk.Frame(frame)
        text_frame.pack(fill="x")
        initial_height = self.settings.get("result_height", 14)
        text_widget = tk.Text(text_frame, height=initial_height, wrap="word", state="disabled")
        scroll = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        text_widget.config(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        text_widget.pack(side="left", fill="both", expand=True)
        self._add_copy_context_menu(text_widget)

        return frame, text_widget

    # --- 対象日付範囲の選択・結果欄の高さを次回起動時のデフォルトとして保存する ---
    def _save_settings(self):
        self.settings["range"] = self.range_var.get()
        if hasattr(self, "result_text"):
            self.settings["result_height"] = int(self.result_text.cget("height"))
        save_settings(self.settings)

    # ---------- 右クリックでコピーできるようにする ----------

    def _add_copy_context_menu(self, text_widget):
        menu = tk.Menu(text_widget, tearoff=0)
        menu.add_command(label="コピー", command=lambda: self._copy_selection(text_widget))
        menu.add_command(label="すべて選択", command=lambda: self._select_all_text(text_widget))

        def show_menu(event):
            menu.tk_popup(event.x_root, event.y_root)

        text_widget.bind("<Button-3>", show_menu)

    def _copy_selection(self, text_widget):
        try:
            selected = text_widget.get("sel.first", "sel.last")
        except tk.TclError:
            return
        self.clipboard_clear()
        self.clipboard_append(selected)

    def _select_all_text(self, text_widget):
        text_widget.tag_add("sel", "1.0", "end")

    def _add_edit_context_menu(self, text_widget):
        menu = tk.Menu(text_widget, tearoff=0)
        menu.add_command(label="切り取り", command=lambda: text_widget.event_generate("<<Cut>>"))
        menu.add_command(label="コピー", command=lambda: text_widget.event_generate("<<Copy>>"))
        menu.add_command(label="貼り付け", command=lambda: text_widget.event_generate("<<Paste>>"))
        menu.add_separator()
        menu.add_command(label="すべて選択", command=lambda: self._select_all_text(text_widget))

        def show_menu(event):
            menu.tk_popup(event.x_root, event.y_root)

        text_widget.bind("<Button-3>", show_menu)

    # ---------- 仕事の進め方ガイド ----------

    def _show_instructions_dialog(self, auto=False):
        if self._instructions_win is not None and self._instructions_win.winfo_exists():
            self._instructions_win.lift()
            self._instructions_win.focus_force()
            return

        win = tk.Toplevel(self)
        self._instructions_win = win
        win.title(f"仕事の進め方({os.path.basename(os.path.dirname(os.path.abspath(__file__)))})")
        win.geometry("560x520")
        win.transient(self)

        text_frame = ttk.Frame(win)
        text_frame.pack(fill="both", expand=True, padx=14, pady=(14, 6))

        text_widget = tk.Text(text_frame, wrap="word", state="normal")
        scroll = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
        text_widget.config(yscrollcommand=scroll.set)
        scroll.pack(side="right", fill="y")
        text_widget.pack(side="left", fill="both", expand=True)

        self._insert_with_links(text_widget, read_instructions())
        text_widget.config(state="disabled")
        self._add_copy_context_menu(text_widget)

        button_row = ttk.Frame(win)
        button_row.pack(side="bottom", fill="x", padx=14, pady=(0, 14))
        ttk.Button(button_row, text="✎ 取説を編集", command=self._open_instructions_editor).pack(side="left")
        ttk.Button(button_row, text="閉じる", command=win.destroy).pack(side="right")

        if not auto:
            win.lift()
            win.focus_force()

    def _open_instructions_editor(self):
        path = ensure_instructions_file()
        import subprocess
        try:
            subprocess.Popen(["notepad.exe", path])
        except OSError as e:
            messagebox.showerror("編集に失敗しました", f"{path}\n{e}")

    def _insert_with_links(self, text_widget, content):
        """Markdownの簡易記法(見出し・引用・コードブロック・太字・区切り線)をそれらしく
        整形しつつ、URL・ファイルパスをクリック可能なリンクとして挿入する。"""
        from tkinter import font as tkfont
        md_inline_pattern = re.compile(
            r"(?P<bold>\*\*(?P<boldtext>.+?)\*\*)"
            rf"|(?P<code>`(?P<codetext>[^`\n]+)`)"
            rf"|(?P<link>https?://{_LINK_CHAR}+|[A-Za-z]:\\{_LINK_CHAR}+|\\\\{_LINK_CHAR}+)"
        )
        try:
            base_font = tkfont.nametofont(text_widget.cget("font"))
            family, size = base_font.actual("family"), base_font.actual("size")
        except Exception:
            family, size = "TkDefaultFont", 10

        text_widget.tag_configure("hyperlink", foreground="#2f6f63", underline=True)
        text_widget.tag_configure("md_h1", font=(family, size + 5, "bold"), spacing1=10, spacing3=6)
        text_widget.tag_configure("md_h2", font=(family, size + 3, "bold"), spacing1=8, spacing3=4)
        text_widget.tag_configure("md_h3", font=(family, size + 1, "bold"), spacing1=6, spacing3=3)
        text_widget.tag_configure("md_bold", font=(family, size, "bold"))
        text_widget.tag_configure("md_code", font=("Consolas", size), background="#eeeeee")
        text_widget.tag_configure("md_code_block", font=("Consolas", size), background="#f2f2f2", lmargin1=14, lmargin2=14)
        text_widget.tag_configure("md_quote", foreground="#4a4a4a", background="#f0efe6", lmargin1=20, lmargin2=20, spacing1=3, spacing3=3)
        text_widget.tag_configure("md_hr", foreground="#aaaaaa")

        orig_cursor = text_widget.cget("cursor")
        link_index = 0

        def insert_inline(text_line, block_tags):
            nonlocal link_index
            pos = 0
            for m in md_inline_pattern.finditer(text_line):
                start, end = m.span()
                if start > pos:
                    text_widget.insert("end", text_line[pos:start], block_tags)
                if m.group("bold"):
                    text_widget.insert("end", m.group("boldtext"), block_tags + ("md_bold",))
                elif m.group("code"):
                    text_widget.insert("end", m.group("codetext"), block_tags + ("md_code",))
                else:
                    raw = m.group("link")
                    href = raw.rstrip(LINK_TRAILING_CHARS)
                    trailing = raw[len(href):]
                    tag_name = f"link_{link_index}"
                    link_index += 1
                    text_widget.insert("end", href, block_tags + ("hyperlink", tag_name))
                    text_widget.tag_bind(tag_name, "<Button-1>", lambda e, href=href: self._open_link(href))
                    text_widget.tag_bind(tag_name, "<Enter>", lambda e: text_widget.config(cursor="hand2"))
                    text_widget.tag_bind(tag_name, "<Leave>", lambda e: text_widget.config(cursor=orig_cursor))
                    if trailing:
                        text_widget.insert("end", trailing, block_tags)
                pos = end
            text_widget.insert("end", text_line[pos:], block_tags)

        in_code_block = False
        for raw_line in content.splitlines():
            stripped = raw_line.strip()

            if stripped.startswith("```"):
                in_code_block = not in_code_block
                continue
            if in_code_block:
                text_widget.insert("end", raw_line + "\n", ("md_code_block",))
                continue

            if re.fullmatch(r"[-_*]{3,}", stripped):
                text_widget.insert("end", "─" * 42 + "\n", ("md_hr",))
                continue

            m = re.match(r"^(#{1,6})\s+(.*)$", raw_line)
            if m:
                level = len(m.group(1))
                tag = "md_h1" if level == 1 else "md_h2" if level == 2 else "md_h3"
                insert_inline(m.group(2), (tag,))
                text_widget.insert("end", "\n")
                continue

            m = re.match(r"^>\s?(.*)$", raw_line)
            if m:
                inner = m.group(1)
                hm = re.match(r"^(#{1,6})\s+(.*)$", inner)
                if hm:
                    level = len(hm.group(1))
                    tag = "md_h1" if level == 1 else "md_h2" if level == 2 else "md_h3"
                    insert_inline(hm.group(2), ("md_quote", tag))
                else:
                    insert_inline(inner, ("md_quote",))
                text_widget.insert("end", "\n")
                continue

            insert_inline(raw_line, ())
            text_widget.insert("end", "\n")

    def _open_link(self, href):
        if href.startswith("http://") or href.startswith("https://"):
            webbrowser.open(href)
            return
        try:
            os.startfile(href)
        except OSError as e:
            messagebox.showerror("開けませんでした", f"{href}\n{e}")

    # ---------- 施設検索 ----------

    def _refresh_store_status(self):
        if self.store_data.error:
            self.store_status_var.set(self.store_data.error)
        else:
            self.store_status_var.set("{}件読み込み済み".format(len(self.store_data.rows)))

    def _reload_store_data(self):
        self.store_data.reload()
        self._refresh_store_status()
        self._on_store_search(None)

    def _on_store_search(self, event):
        keyword = self.store_entry.get().strip()
        self.store_result_text.config(state="normal")
        self.store_result_text.delete("1.0", "end")
        if self.store_data.error:
            pass
        elif keyword:
            results = self.store_data.search(keyword)
            if results:
                lines = ["{}  →  {}".format(name, area) for name, area in results[:200]]
                self.store_result_text.insert("end", "\n".join(lines))
            else:
                self.store_result_text.insert("end", "該当なし")
        self.store_result_text.config(state="disabled")

    def _on_cal_text_paste(self, event):
        # <<Paste>>はテキスト挿入前に発火することがあるため、挿入完了後に処理する
        self.after(10, self._auto_fill_store_from_cal_text)

    def _auto_fill_store_from_cal_text(self):
        text = self.cal_text.get("1.0", "end")
        facility = extract_facility_name(text)
        if facility:
            self.store_entry.delete(0, "end")
            self.store_entry.insert(0, facility)
            self._on_store_search(None)

    # ---------- イベント ----------

    def _on_pdf_drop(self, event):
        path = event.data.strip("{}")  # Windowsでスペースを含むパスは{}で囲まれる
        self._set_pdf(path)

    def _choose_pdf(self):
        path = filedialog.askopenfilename(title="日報PDF(欠席時対応記録票)を開く", filetypes=[("PDFファイル", "*.pdf")])
        if path:
            self._set_pdf(path)

    def _open_pdf(self):
        if not self.pdf_path:
            messagebox.showerror("エラー", "先にPDFを選択してください。")
            return
        os.startfile(self.pdf_path)

    def _set_pdf(self, path):
        if not path.lower().endswith(".pdf"):
            messagebox.showerror("エラー", "PDFファイルを指定してください。")
            return
        self.pdf_path = path
        self.pdf_label.config(text=f"選択済み: {path}")

        # 新しいPDFが来たら、前回分の貼り付けテキストと結果表示をクリアして
        # 新しい情報の受付待ち状態に戻す(年月欄は「今」で埋め直す。対象日付範囲は維持する)
        self.cal_text.delete("1.0", "end")
        self.year_entry.delete(0, "end")
        self.year_entry.insert(0, str(datetime.now().year))
        self.month_entry.delete(0, "end")
        self.month_entry.insert(0, str(datetime.now().month))
        self.result_text.config(state="normal")
        self.result_text.delete("1.0", "end")
        self.result_text.config(state="disabled")
        self.blank_text.config(state="normal")
        self.blank_text.delete("1.0", "end")
        self.blank_text.config(state="disabled")

    def _run_compare(self):
        if not self.pdf_path:
            messagebox.showwarning("確認", "PDFがまだ選択されていません。")
            return

        cal_text = self.cal_text.get("1.0", "end").strip()
        if not cal_text:
            messagebox.showwarning("確認", "カレンダーのテキストが貼り付けられていません。")
            return

        year_str = self.year_entry.get().strip()
        month_str = self.month_entry.get().strip()
        year = int(year_str) if year_str.isdigit() else None
        month = int(month_str) if month_str.isdigit() else None

        try:
            pdf_rows = extract_pdf_rows(self.pdf_path)
            date_errors = find_pdf_date_errors(self.pdf_path)
            pdf_flag_map = get_pdf_flag_map(self.pdf_path)
            pdf_duplicates = find_pdf_duplicates(self.pdf_path)
            pdf_blank_flag_rows = find_pdf_blank_flag_rows(self.pdf_path)
        except Exception as e:
            messagebox.showerror("PDF読み込みエラー", str(e))
            return

        cal_rows_all = parse_calendar_text(cal_text, month)
        cal_rows = [r for r in cal_rows_all if r.get("status") == "○"]
        cal_nokasan_rows = [r for r in cal_rows_all if r.get("status") == "✕"]

        if not pdf_rows:
            messagebox.showwarning("確認", "PDFから「欠席加算=○」の行が見つかりませんでした。フォーマットをご確認ください。")
        if not cal_rows_all:
            messagebox.showwarning("確認", "貼り付けたテキストから「欠席Ⅰ」が見つかりませんでした。コピー範囲・書式をご確認ください。")

        # 対象日付範囲(1日〜15日 / 16日〜25日 / 26日〜月末)で絞り込む。月末は年月から算出する。
        range_choice = self.range_var.get()
        if year and month:
            last_day = calendar_mod.monthrange(year, month)[1]
        else:
            last_day = 31  # 年月が未入力/不正な場合は安全側(31日まで)で判定

        if range_choice == "first_third":
            range_start, range_end = 1, 15
        elif range_choice == "second_third":
            range_start, range_end = 16, 25
        else:
            range_start, range_end = 26, last_day

        def in_range(date_str):
            day = extract_day(date_str)
            if day is None:
                return True  # 日にちが読み取れない行はフィルタせずそのまま残す
            return range_start <= day <= range_end

        pdf_rows = [r for r in pdf_rows if in_range(r["date"])]
        cal_rows = [r for r in cal_rows if in_range(r["date"])]
        cal_nokasan_rows = [r for r in cal_nokasan_rows if in_range(r["date"])]
        # 日付誤記疑い(受付日＞欠席日)・重複の疑いは対象範囲に関わらず全件チェックする
        pdf_blank_flag_rows = [r for r in pdf_blank_flag_rows if in_range(r["date"])]

        matched, only_pdf, only_cal = compare_rows(pdf_rows, cal_rows)

        # HUGで「欠席Ⅰ（加算なし)」(=✕)となっている人のキー(日付+氏名)。
        # only_pdf(PDFには〇があるがHUGの「欠席Ⅰ」には無い人)がこれに該当する場合、
        # 「HUGに記載が無い」のではなく「HUGでは✕(加算なし)扱い」なので表示を差し替える
        nokasan_keys = {
            (normalize_date(r["date"]), normalize_name(r["name"])) for r in cal_nokasan_rows
        }

        # 「HUG(カレンダー)」「日報(PDF)」それぞれに欠席Ⅰ/欠席加算〇があるかを
        # date・氏名ごとにまとめ、日付順に並べる(日報側にPDFの行番号・空欄有無があれば付ける)
        merged = (
            [
                (
                    "欠席チェック不一致",
                    format_mmdd(r["date"], month),
                    r["name"],
                    "✕" if (normalize_date(r["date"]), normalize_name(r["name"])) in nokasan_keys else "記載なし",
                    "〇",
                    r.get("row_no", ""),
                    "",
                    "あり" if r.get("has_blank") else "",
                )
                for r in only_pdf
            ]
            + [
                (
                    "欠席チェック不一致",
                    format_mmdd(r["date"], month),
                    r["name"],
                    "〇",
                    pdf_info["flag"] if pdf_info else "記載なし",
                    pdf_info["row_no"] if pdf_info else "",
                    "",
                    "",
                )
                for r in only_cal
                for pdf_info in [pdf_flag_map.get((normalize_date(r["date"]), normalize_name(r["name"])))]
            ]
        )
        merged.sort(key=lambda x: (x[1], x[2]))

        # 欠席日 < 受付日 になっている行(日付の誤記の疑い)
        date_error_rows = [
            (
                "日付誤記疑い",
                format_mmdd(r["date"], month),
                r["name"],
                "-",
                "-",
                r["row_no"],
                format_mmdd(r["uketsuke"], month),
                "",
            )
            for r in date_errors
        ]
        date_error_rows.sort(key=lambda x: (x[1], x[2]))

        # 一致しているが、日報側の〇の行に空欄がある(記入漏れの疑い)ものを別枠でまとめる
        matched_blank_rows = [
            (r["name"], r.get("row_no", ""), format_mmdd(r["date"], month))
            for r in matched
            if r.get("has_blank")
        ]
        matched_blank_rows.sort(key=lambda x: (x[2], x[0]))

        # 不一致(only_pdf)だが、日報側の〇の行に空欄がある(記入漏れの疑い)ものも右側にまとめる
        only_pdf_blank_rows = [
            (r["name"], r.get("row_no", ""), format_mmdd(r["date"], month))
            for r in only_pdf
            if r.get("has_blank")
        ]
        only_pdf_blank_rows.sort(key=lambda x: (x[2], x[0]))

        pdf_duplicates.sort(key=lambda d: (normalize_date(d["date"]), d["name"]))
        pdf_blank_flag_rows.sort(key=lambda r: (normalize_date(r["date"]), r["name"]))

        range_label = f"{range_start}日〜{range_end}日"
        self._show_result(
            len(matched),
            merged,
            date_error_rows,
            matched_blank_rows,
            only_pdf_blank_rows,
            pdf_duplicates,
            pdf_blank_flag_rows,
            month,
            range_label,
        )

    def _show_result(
        self,
        matched_count,
        merged,
        date_error_rows,
        matched_blank_rows,
        only_pdf_blank_rows,
        pdf_duplicates,
        pdf_blank_flag_rows,
        month,
        range_label="",
    ):
        self.result_text.config(state="normal")
        self.result_text.delete("1.0", "end")

        lines = [f"■ 対象範囲: {range_label}", f"■ HUGと日報の一致: {matched_count}件", ""]
        lines.append(f"■ 不一致(HUGと日報で欠席Ⅰ/欠席加算〇が食い違っている): {len(merged)}件")
        for _, date, name, hug, nippo, row_no, _uke, _has_blank in merged:
            row_no_disp = row_no if row_no else "-"
            lines.append(f"{name}\t{row_no_disp}\t{date}\tHUG:{hug}、日報:{nippo}")

        lines.append("")
        lines.append(f"■ 日付の誤記の疑い(〇の行のみ対象・受付日が欠席日より後になっている): {len(date_error_rows)}件")
        for _, date, name, _hug, _nippo, row_no, uketsuke, _blank in date_error_rows:
            lines.append(f"{name}\t{row_no}\t{date}\t受付日:{uketsuke}　欠席日:{date}")

        lines.append("")
        lines.append(f"■ 重複の疑い(氏名+欠席日がPDF内で完全一致): {len(pdf_duplicates)}件")
        for d in pdf_duplicates:
            row_no_str = "と".join((f"{rn}番" if rn else "-") for rn in d["row_nos"])
            lines.append(f"{d['name']}\t行番号{row_no_str}　{format_mmdd(d['date'], month)}\t重複")

        self.result_text.insert("1.0", "\n".join(lines))
        self.result_text.config(state="disabled")

        # 右側: 一致しているが〇の行に空欄がある(記入漏れの疑い)
        self.blank_text.config(state="normal")
        self.blank_text.delete("1.0", "end")
        blank_lines = [f"■ 該当: {len(matched_blank_rows)}件", ""]
        for name, row_no, date in matched_blank_rows:
            row_no_disp = row_no if row_no else "-"
            blank_lines.append(f"{name}\t{row_no_disp}\t{date}\t日報:空欄あり")

        blank_lines.append("")
        blank_lines.append(f"■ 不一致だが〇の行に空欄がある(記入漏れの疑い): {len(only_pdf_blank_rows)}件")
        for name, row_no, date in only_pdf_blank_rows:
            row_no_disp = row_no if row_no else "-"
            blank_lines.append(f"{name}\t{row_no_disp}\t{date}\t日報:空欄あり")

        blank_lines.append("")
        blank_lines.append(f"■ 欠席加算欄が未記入(空欄): {len(pdf_blank_flag_rows)}件")
        for r in pdf_blank_flag_rows:
            row_no_disp = r["row_no"] if r["row_no"] else "-"
            blank_lines.append(f"{r['name']}\t{row_no_disp}\t{format_mmdd(r['date'], month)}")

        self.blank_text.insert("1.0", "\n".join(blank_lines))
        self.blank_text.config(state="disabled")

if __name__ == "__main__":
    if sys.platform == "win32":
        import ctypes

        console_window = ctypes.windll.kernel32.GetConsoleWindow()
        if console_window:
            ctypes.windll.user32.ShowWindow(console_window, 6)  # SW_MINIMIZE

    if not DND_AVAILABLE:
        print("※ tkinterdnd2が見つからないため、ドラッグ&ドロップは使えません。")
        print("  pip install tkinterdnd2 を実行すると有効になります。")
    app = App()
    app.mainloop()
